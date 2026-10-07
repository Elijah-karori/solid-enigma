#!/usr/bin/env python3
"""
Extract EVERYTHING from the Apps Script inventory workbook into a clean, validated
JSON bundle (import_bundle.json) plus a human-readable reconciliation report.

The Go loader (cmd/import) reads the bundle; this script never touches the database
and never exports passwords (legacy hashes / plaintext are dropped on purpose).

Usage:
  python3 tools/extract_workbook.py path/to/workbook.xlsx [--out out/] [--as-of YYYY-MM-DD]

Requires: pip install openpyxl
"""
import argparse, collections, datetime as dt, hashlib, json, os, re, sys
import openpyxl

PLACEHOLDERS = {"", "n/a", "na", "none", "null", "-", "not stated", "not legible", "unknown", "nil"}
SERIAL_PREFIXES = ("SKU-ONT", "SKU-RTR", "SKU-NET", "SKU-FAT")
STATUS_MAP = {
    "in stock": "In Stock", "stock in": "In Stock",
    "issued / out": "Issued / Out", "issued": "Issued / Out", "stock out": "Issued / Out",
    "out": "Issued / Out", "dispatched": "Issued / Out",
    "under repair": "Under Repair", "decommissioned": "Decommissioned",
}
ROLE_MAP = {"store supervisor": "Store Manager"}
ROLES = ["Admin", "Store Manager", "Finance", "Project Manager", "Support", "Technician"]
FS = "\u241f"  # field separator used by the legacy audit hash

issues = collections.defaultdict(list)  # category -> [messages]


def warn(cat, msg):
    issues[cat].append(msg)


def s(v):
    """Trimmed string; real blank -> ''."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def clean(v):
    """String with placeholder text ('Not legible', 'N/A', ...) treated as missing."""
    t = s(v)
    return "" if t.lower() in PLACEHOLDERS else t


def num(v, default=0.0):
    try:
        if v is None or v == "":
            return default
        return float(v)
    except (TypeError, ValueError):
        return default


def to_dt(v):
    """datetime | Excel serial | ISO-ish string -> datetime or None."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and 20000 < float(v) < 80000:  # Excel serial day number
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))
    t = s(v)
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(t, fmt)
        except ValueError:
            pass
    return None


def iso(v):
    d = to_dt(v)
    return d.strftime("%Y-%m-%dT%H:%M:%S") if d else ""


def day(v):
    d = to_dt(v)
    return d.strftime("%Y-%m-%d") if d else ""


def norm_mac(raw, ctx):
    t = clean(raw)
    if not t:
        return "", ""
    flat = re.sub(r"[^0-9A-Fa-f]", "", t)
    if re.fullmatch(r"[0-9A-Fa-f]{12}", re.sub(r"[:.\-\s]", "", t)):
        return re.sub(r"[:.\-\s]", "", t).upper(), ""
    head = re.match(r"^([0-9A-Fa-f]{2}[:.\-]?){5}[0-9A-Fa-f]{2}", t)
    if head:
        m = re.sub(r"[:.\-]", "", head.group(0)).upper()
        warn("serialized", f"{ctx}: MAC '{t}' had trailing text; kept {m}")
        return m, f"Original MAC text: {t}"
    warn("serialized", f"{ctx}: MAC '{t}' is not 12 hex characters; stored blank (kept in notes)")
    return "", f"Original MAC text: {t}"


def sheet_rows(wb, name, width):
    """Rows with at least one meaningful cell in the first `width` columns (skips formula filler)."""
    if name not in wb.sheetnames:
        warn("workbook", f"sheet '{name}' not found")
        return []
    out = []
    for r in wb[name].iter_rows(min_row=2, values_only=True):
        r = list(r[:width]) + [None] * (width - len(r[:width]))
        if any(c not in (None, "") for c in r):
            out.append(r)
    return out


def sku_kind(sku, asset_type=""):
    if sku.upper().startswith(SERIAL_PREFIXES):
        return "Serialized"
    return "Bulk"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default="out")
    ap.add_argument("--as-of", default="", help="date used for opening-balance adjustments (default: last ledger date)")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    counts = {}

    # ------------------------------------------------------------- settings
    settings = []
    for k, v in sheet_rows(wb, "Inventory Settings", 2):
        val = "TRUE" if v is True else "FALSE" if v is False else s(v)
        settings.append({"key": s(k), "value": val})
    counts["settings"] = len(settings)

    # ---------------------------------------------------------------- users
    users, by_email, by_name = [], {}, collections.defaultdict(list)
    for r in sheet_rows(wb, "User Directory", 10):
        email = s(r[0]).lower()
        if not email:
            continue
        role = ROLE_MAP.get(s(r[2]).lower(), s(r[2]))
        match = next((x for x in ROLES if x.lower() == role.lower()), None)
        if not match:
            warn("users", f"{email}: unknown role '{role}' -> imported as Technician (least privilege)")
            match = "Technician"
        pw = r[1]
        kind = "none" if pw in (None, "") else "hashed" if s(pw).startswith("sha256$") else "plaintext"
        if kind == "plaintext":
            warn("users", f"{email}: legacy PLAINTEXT password found in sheet - discarded, forced reset")
        u = {
            "legacy_id": s(r[6]), "email": email, "name": s(r[3]) or email, "role": match,
            "status": "inactive" if s(r[7]).lower() in ("inactive", "disabled", "suspended") else "active",
            "site_station": s(r[8]), "contact_info": s(r[9]),
            "must_reset_password": True,  # no credential is ever carried over
            "legacy_password_kind": kind,
        }
        users.append(u)
        by_email[email] = u
        by_name[u["name"].lower()].append(u)
        by_name[u["name"].split()[0].lower()].append(u)
    counts["users"] = len(users)

    def resolve_user(text):
        """Name or email from a free-text cell -> (name, email)."""
        t = s(text)
        if not t:
            return "", ""
        if t.lower() in by_email:
            u = by_email[t.lower()]
            return u["name"], u["email"]
        hit = {x["email"]: x for x in by_name.get(t.lower(), [])}
        if len(hit) == 1:
            u = next(iter(hit.values()))
            return u["name"], u["email"]
        return t, ""

    # -------------------------------------------------------------- catalog
    catalog = {}
    for r in sheet_rows(wb, "Item Catalog", 10):
        sku = s(r[0]).upper()
        if not sku or sku in catalog:
            if sku:
                warn("catalog", f"duplicate SKU {sku} ignored")
            continue
        tracking = s(r[8]) or sku_kind(sku)
        trk = {"Serialized": "SERIALIZED", "Non-Serialized": "NONSER", "Bulk": "BULK"}.get(tracking, "BULK")
        manu = clean(r[2])
        if re.fullmatch(r"[0-9A-Fa-f]{12}", manu or ""):
            warn("catalog", f"{sku}: manufacturer looked like a MAC ({manu}); cleared")
            manu = ""
        catalog[sku] = {
            "sku": sku, "asset_type": s(r[1]), "manufacturer": manu, "model": s(r[3]),
            "access_tech": clean(r[4]), "unit_cost": num(r[5]), "description": s(r[6]),
            "reorder_level": int(num(r[7], 3)), "tracking_type": trk,
            "project_scope": "" if s(r[9]).lower() in ("", "shared") else s(r[9]).upper(),
            "is_active": True,
        }
    counts["catalog"] = len(catalog)

    # ----------------------------------------------------------- serialized
    serialized, sn_seen, mac_seen = [], {}, {}
    model_to_sku = {}
    for c in catalog.values():
        model_to_sku.setdefault(c["model"].lower(), []).append(c["sku"])
    prefix_by_asset = {"INV-ONT": "SKU-ONT", "INV-RTR": "SKU-RTR", "INV-NET": "SKU-NET", "INV-FAT": "SKU-FAT"}

    for r in sheet_rows(wb, "Serialized Inventory", 17):
        aid = s(r[0]).upper()
        if not aid:
            continue
        notes = [s(r[12])] if s(r[12]) else []
        sku = s(r[13]).upper()
        model = s(r[3])
        if not sku:
            hits = [k for m, ks in model_to_sku.items() if m and (m in model.lower() or model.lower().startswith(m)) for k in ks]
            hits = [h for h in hits if h.startswith(prefix_by_asset.get(aid[:7], "SKU-"))]
            if len(set(hits)) == 1:
                sku = hits[0]
                notes.append(f"SKU inferred from model '{model}'")
                warn("serialized", f"{aid}: SKU blank, inferred {sku} from model")
            else:
                warn("serialized", f"{aid}: SKU blank and model '{model}' matches no catalog item - NEEDS MANUAL SKU")
        mac, mnote = norm_mac(r[6], aid)
        if mnote:
            notes.append(mnote)
        sn_raw = s(r[7])
        sn = clean(r[7]).upper().replace(" ", "")
        if sn_raw and not sn:
            notes.append(f"Serial recorded as '{sn_raw}' in source")
            warn("serialized", f"{aid}: serial '{sn_raw}' is not a real value -> stored NULL")
        elif not sn_raw:
            warn("serialized", f"{aid}: serial number blank -> stored NULL")
        if sn:
            if sn in sn_seen:
                warn("serialized", f"{aid}: serial {sn} duplicates {sn_seen[sn]} -> stored NULL on this unit")
                notes.append(f"Duplicate serial {sn} (also {sn_seen[sn]}) in source")
                sn = ""
            else:
                sn_seen[sn] = aid
        if mac:
            if mac in mac_seen:
                warn("serialized", f"{aid}: MAC {mac} duplicates {mac_seen[mac]}")
            else:
                mac_seen[mac] = aid
        raw_status = s(r[8])
        status = STATUS_MAP.get(raw_status.lower())
        if not status:
            warn("serialized", f"{aid}: unknown status '{raw_status}' -> In Stock")
            status = "In Stock"
        if raw_status != status:
            warn("status-normalised", f"{aid}: '{raw_status}' -> '{status}'")
        cond = s(r[9]) or "Not Recorded"
        serialized.append({
            "asset_id": aid, "sku": sku, "asset_type": s(r[1]), "manufacturer": clean(r[2]),
            "model": clean(r[3]) or model, "access_tech": clean(r[4]), "product_id": clean(r[5]),
            "mac": mac, "serial_number": sn or None, "status": status, "condition": cond,
            "location": s(r[10]), "custodian": s(r[11]), "notes": " | ".join(n for n in notes if n),
            "customer_name": s(r[14]), "customer_account": s(r[15]), "linked_ticket_id": s(r[16]),
        })
    counts["serialized"] = len(serialized)

    # --------------------------------------------------------------- ledger
    ledger_rows = sheet_rows(wb, "Transaction Ledger", 15)
    txns, bal = [], collections.Counter()
    for r in ledger_rows:
        d = to_dt(r[0])
        direction = {"stock in": "Stock In", "stock out": "Stock Out"}.get(s(r[1]).lower(), "")
        sku = s(r[2]).upper()
        qty = num(r[4])
        if not (d and direction and sku and qty > 0):
            warn("ledger", f"row skipped (date/direction/sku/qty invalid): {r[:6]}")
            continue
        # unit/total cost in the workbook are formulas; recompute from catalog when cached value is missing
        unit = num(r[8], catalog.get(sku, {}).get("unit_cost", 0.0))
        total = num(r[9], unit * qty)
        txns.append({
            "transaction_date": d.strftime("%Y-%m-%dT%H:%M:%S"), "direction": direction, "sku": sku,
            "item_name": s(r[3]), "quantity": qty, "requested_by": s(r[5]), "role": s(r[6]),
            "site_reference": s(r[7]), "unit_cost": unit, "total_cost": total,
            "cost_type": s(r[10]) or ("Procurement / Stock In" if direction == "Stock In" else "Other"),
            "task_id": s(r[11]), "asset_id": s(r[12]), "project_id": s(r[14]), "notes": "Imported from Transaction Ledger",
        })
        bal[sku] += qty if direction == "Stock In" else -qty
        if sku not in catalog:
            warn("ledger", f"ledger SKU {sku} is not in the Item Catalog")
    counts["transactions_ledger"] = len(txns)
    as_of = a.as_of or max((t["transaction_date"][:10] for t in txns), default=dt.date.today().isoformat())

    # ---------------------------------------- stock summary (cached) vs ledger
    cached = {s(r[0]).upper(): int(num(r[4])) for r in sheet_rows(wb, "Stock Summary", 6) if s(r[0])}
    stale = []
    for sku in sorted(set(cached) | set(bal)):
        if int(bal.get(sku, 0)) != cached.get(sku, 0):
            stale.append((sku, cached.get(sku), int(bal.get(sku, 0))))
            warn("summary-stale", f"{sku}: Stock Summary cached {cached.get(sku)} but ledger replays to {int(bal.get(sku, 0))} (ledger wins)")

    # ---------------- opening-balance adjustments for serialized SKUs --------
    # Where unit records exist, the units are the physical truth. Ledger drift (for example the
    # 'Backfilled from Serialized Inventory' rows that were posted as Stock Out) is corrected by
    # explicit, labelled Adjustment rows so history stays intact and stock never starts negative.
    units_in_stock = collections.Counter(u["sku"] for u in serialized if u["status"] == "In Stock" and u["sku"])
    skus_with_units = {u["sku"] for u in serialized if u["sku"]}
    adjustments = []
    for sku in sorted(skus_with_units):
        diff = units_in_stock.get(sku, 0) - int(bal.get(sku, 0))
        if diff:
            c = catalog.get(sku, {})
            adjustments.append({
                "transaction_date": f"{as_of}T00:00:00", "direction": "Stock In" if diff > 0 else "Stock Out", "sku": sku,
                "item_name": c.get("model", ""), "quantity": abs(diff), "requested_by": "Migration", "role": "System",
                "site_reference": "Opening balance", "unit_cost": 0.0, "total_cost": 0.0, "cost_type": "Adjustment",
                "task_id": "", "asset_id": "", "project_id": "",
                "notes": f"Migration adjustment: {units_in_stock.get(sku, 0)} unit(s) In Stock in Serialized Inventory vs ledger balance {int(bal.get(sku, 0))}",
            })
            bal[sku] += diff
    # serialized SKUs with ledger balance but no unit records -> flag, do not invent units
    for sku in sorted(c for c in catalog if catalog[c]["tracking_type"] == "SERIALIZED" and c not in skus_with_units):
        if bal.get(sku, 0) != 0:
            warn("serialized", f"{sku}: ledger balance {int(bal[sku])} but NO unit records - add S/N and MAC via serialized Stock In/edit")
    counts["transactions_adjustments"] = len(adjustments)

    still_negative = {k: int(v) for k, v in bal.items() if v < 0}
    for k, v in still_negative.items():
        warn("negative", f"{k}: balance {v} after adjustments - review manually")

    # ---------------------------------------- horizontal register classification
    ledger_multiset = collections.Counter((t["direction"], t["sku"], t["quantity"]) for t in txns)
    register_flags = []
    for r in sheet_rows(wb, "Horizontal Register", 12):
        direction = {"stock in": "Stock In", "stock out": "Stock Out"}.get(s(r[1]).lower(), "")
        item = s(r[2]).upper()
        qty = num(r[4])
        key = (direction, item, qty)
        flip = ("Stock In" if direction == "Stock Out" else "Stock Out", item, qty)
        if ledger_multiset.get(key, 0) > 0:
            ledger_multiset[key] -= 1
            continue
        kind = "direction-suspect" if ledger_multiset.get(flip, 0) > 0 and "received" in s(r[7]).lower() else "no-ledger-match"
        register_flags.append({"date": day(r[0]), "direction": direction, "item": item, "qty": qty, "site": s(r[7]), "notes": s(r[8]), "classification": kind})
    counts["register_rows_without_ledger_match"] = len(register_flags)

    # ------------------------------------------------------------ requisitions
    reqs = []
    for r in sheet_rows(wb, "Technician Requisitions", 17):
        if not s(r[0]):
            continue
        tech, _ = resolve_user(r[2])
        reqs.append({
            "requisition_id": s(r[0]).upper(), "date_requested": iso(r[1]), "technician_name": tech,
            "item_sku": s(r[3]).upper(), "quantity_requested": num(r[4]), "reason_job_ticket": s(r[5]),
            "approval_status": s(r[6]), "approved_by": s(r[7]), "issue_date": iso(r[8]), "issued_by": s(r[9]),
            "decision_note": s(r[10]), "project_id": s(r[11]).upper(), "finance_status": s(r[12]) or "N/A",
            "est_value": num(r[13]), "issued_units": s(r[14]),
        })
    counts["requisitions"] = len(reqs)

    # ------------------------------------------------------------------ tasks
    tasks = []
    for r in sheet_rows(wb, "Tasks", 17):
        if not s(r[0]):
            continue
        name, email = resolve_user(r[4])
        tasks.append({
            "task_id": s(r[0]).upper(), "created_at": iso(r[1]), "task_title": s(r[2]), "task_type": s(r[3]),
            "assigned_personnel": name, "assignee_email": s(r[5]) or email, "priority": s(r[6]) or "Normal", "status": s(r[7]),
            "required_sku": s(r[8]).upper(), "required_qty": num(r[9]), "customer_site": s(r[10]), "reference": s(r[11]),
            "notes": s(r[12]), "created_by": s(r[13]), "stock_ready_at": iso(r[14]), "linked_ticket_id": s(r[15]), "project_id": s(r[16]).upper(),
        })
    counts["tasks"] = len(tasks)

    # ---------------------------------------------------------------- tickets
    tickets = []
    for r in sheet_rows(wb, "Customer Issues", 16):
        if not s(r[0]):
            continue
        tech, _ = resolve_user(r[4])
        tickets.append({
            "ticket_id": s(r[0]).upper(), "date_logged": iso(r[1]), "customer_account": s(r[2]), "issue_category": s(r[3]),
            "assigned_technician": tech, "device_swapped_old_sn": clean(r[5]), "replacement_device_new_sn": clean(r[6]),
            "ticket_status": s(r[7]) or "Open", "logged_by": s(r[8]), "resolution_notes": s(r[9]), "linked_task_id": s(r[10]),
            "priority": s(r[11]) or "Normal", "closed_at": iso(r[12]), "customer_id": s(r[13]), "hotspot_id": s(r[14]), "device_asset_id": s(r[15]).upper(),
        })
    counts["tickets"] = len(tickets)

    # ------------------------------------------------------------ procurement
    procs = []
    for r in sheet_rows(wb, "Procurement", 17):
        if not s(r[0]):
            continue
        by, _ = resolve_user(r[2])
        fin, _ = resolve_user(r[10])
        qty, unit = num(r[4]), num(r[5])
        procs.append({
            "procurement_id": s(r[0]).upper(), "date_requested": iso(r[1]), "requested_by": by, "item_sku": s(r[3]).upper(),
            "quantity": qty, "est_unit_cost": unit, "est_total": num(r[6], qty * unit), "project_id": s(r[7]).upper(),
            "linked_requisition_id": s(r[8]), "status": s(r[9]), "finance_by": fin, "finance_date": iso(r[11]),
            "supplier": s(r[12]), "po_ref": s(r[13]), "ordered_date": day(r[14]), "received_date": day(r[15]), "notes": s(r[16]),
        })
        if unit == 0 and s(r[3]):
            warn("procurement", f"{s(r[0])}: estimated unit cost is 0 (catalog has no price for {s(r[3])})")
    counts["procurement"] = len(procs)

    # -------------------------------------------- delivery notes / projects etc.
    notes = [{
        "doc_no": s(r[0]), "type": s(r[1]), "reference": s(r[2]), "project_id": s(r[3]).upper(), "recipient": s(r[4]), "file_name": s(r[5]),
        "drive_url": s(r[6]), "created_by": s(r[7]), "created_at": iso(r[8]), "value_kes": num(r[10]), "payment_status": s(r[11]),
        "paid_date": day(r[12]), "payment_ref": s(r[13]),
    } for r in sheet_rows(wb, "Delivery Notes", 14) if s(r[0])]
    projects = [{
        "project_id": s(r[0]).upper(), "project_name": s(r[1]), "type": s(r[2]), "status": s(r[4]) or "Planning", "start_date": day(r[5]), "end_date": day(r[6]),
        "location_fat": s(r[7]), "project_manager": s(r[8]), "budget": num(r[9]), "notes": s(r[10]), "created_by": s(r[11]),
    } for r in sheet_rows(wb, "Projects & Infrastructure", 12) if s(r[0])]
    customers = [list(r) for r in sheet_rows(wb, "Customers", 20) if s(r[0])]
    hotspots = [list(r) for r in sheet_rows(wb, "Hotspots", 16) if s(r[0])]
    splitters = [list(r) for r in sheet_rows(wb, "Splitters", 11) if s(r[0])]
    enclosures = [list(r) for r in sheet_rows(wb, "Enclosures", 9) if s(r[0])]
    topology = [list(r) for r in sheet_rows(wb, "Network Topology & Locations", 6) if s(r[0])]
    for name, rows in (("Customers", customers), ("Hotspots", hotspots), ("Splitters", splitters), ("Enclosures", enclosures),
                       ("Network Topology & Locations", topology), ("Delivery Notes", notes), ("Projects & Infrastructure", projects)):
        counts[name.lower().replace(" & ", "_").replace(" ", "_")] = len(rows)
        if not rows:
            warn("empty-sheets", f"'{name}' has no data rows - nothing to import")
    if customers or hotspots or splitters or enclosures or topology:
        warn("not-mapped", "Customer/Hotspot/Network rows exist but this extractor writes them raw (bundle keys customers_raw etc.) - map them before import")

    # ---------------------------------------------------------- notification log
    notif = []
    for r in sheet_rows(wb, "Notification Log", 7):
        notif.append({"timestamp": iso(r[0]), "type": s(r[1]), "reference": s(r[2]), "recipient": s(r[3]), "subject": s(r[4]),
                      "status": s(r[5]), "error": s(r[6])[:300], "legacy": True})
    counts["notification_log"] = len(notif)

    # --------------------------------------------- legacy audit trail + chain check
    audit, prev, intact, chained = [], "", True, 0
    for r in sheet_rows(wb, "Audit Trail Log", 11):
        f = [s(c) for c in r[:10]]
        h = s(r[10])
        ok = None
        if h:
            expect = hashlib.sha256((prev + FS + FS.join(f)).encode("utf-8")).hexdigest()
            ok = expect == h
            if not ok:
                intact = False
                warn("audit", f"{f[0]}: stored hash cannot be reproduced from the exported cell values (not necessarily tampering - Sheets may have re-typed a cell). Imported as legacy, unverified.")
            prev, chained = h, chained + 1
        audit.append({"audit_id": f[0], "timestamp": f[1], "entity_type": f[2], "entity_id": f[3], "action": f[4], "user": f[5], "role": f[6],
                      "previous_state": f[7], "new_state": f[8], "details": f[9], "hash": h, "chain_ok": ok})
    counts["audit_legacy"] = len(audit)

    # -------------------------------------------------------------- magic tokens
    live = sum(1 for r in sheet_rows(wb, "Magic Tokens", 3) if (to_dt(r[2]) or dt.datetime.min) > dt.datetime.now())
    counts["magic_tokens_total"] = len(sheet_rows(wb, "Magic Tokens", 3))
    warn("magic-tokens", f"{counts['magic_tokens_total']} token(s) found, {live} unexpired - none imported (replaced by email OTP login)")

    # ------------------------------------------------------------------- output
    bundle = {
        "meta": {"source_file": os.path.basename(a.xlsx), "generated_at": dt.datetime.now().isoformat(timespec="seconds"), "as_of": as_of,
                 "counts": counts, "legacy_audit_chain_intact": intact, "legacy_audit_head_hash": prev},
        "settings": settings, "users": users, "catalog": list(catalog.values()), "serialized": serialized,
        "transactions": txns + adjustments, "requisitions": reqs, "tasks": tasks, "tickets": tickets, "procurement": procs,
        "delivery_notes": notes, "projects": projects, "notification_log": notif, "audit_legacy": audit,
        "expected_balances": {k: int(v) for k, v in sorted(bal.items())},
        "customers_raw": customers, "hotspots_raw": hotspots, "splitters_raw": splitters, "enclosures_raw": enclosures, "topology_raw": topology,
    }
    with open(os.path.join(a.out, "import_bundle.json"), "w", encoding="utf-8") as fh:
        json.dump(bundle, fh, indent=1, ensure_ascii=False)

    # ---------------------------------------------------------------- report
    L = ["# Workbook migration report", "", f"Source: `{os.path.basename(a.xlsx)}`  |  as-of: {as_of}", "", "## Row counts extracted", ""]
    L += [f"- {k}: **{v}**" for k, v in counts.items()]
    L += ["", "## Final expected stock balances (ledger + adjustments)", "", "| SKU | Item | Balance |", "|---|---|---|"]
    for k, v in sorted(bal.items()):
        if v or k in catalog:
            L.append(f"| {k} | {catalog.get(k, {}).get('model', '')} | {int(v)} |")
    L += ["", "## Adjustments generated", ""]
    L += [f"- {x['sku']}: Stock {'In' if x['direction']=='Stock In' else 'Out'} {int(x['quantity'])} - {x['notes']}" for x in adjustments] or ["- none"]
    L += ["", "## Cached Stock Summary disagrees with the ledger", ""]
    L += [f"- {a_}: summary {b_}, ledger {c_}" for a_, b_, c_ in stale] or ["- none"]
    L += ["", "## Horizontal Register rows with no ledger counterpart", ""]
    L += [f"- {x['date']} {x['direction']} {x['item']} x{int(x['qty'])} ({x['site']}) -> {x['classification']}" for x in register_flags] or ["- none"]
    L += ["", f"## Legacy audit chain: {'verified' if intact else 'NOT verifiable from this export'} ({chained} hashed rows). Legacy rows are imported as read-only history; the new chain starts with a genesis entry.", ""]
    for cat, msgs in issues.items():
        L += [f"## Issues: {cat} ({len(msgs)})", ""] + [f"- {m}" for m in msgs] + [""]
    with open(os.path.join(a.out, "migration_report.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))

    print(json.dumps(counts, indent=1))
    print(f"\nWrote {a.out}/import_bundle.json and {a.out}/migration_report.md")
    print("negative balances after adjustments:", still_negative or "none")
    return 0


if __name__ == "__main__":
    sys.exit(main())
