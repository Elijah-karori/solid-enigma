import sqlite3
import openpyxl
import os
import uuid
from datetime import datetime, timedelta

DB_PATH = "app_inventory.db"
EXCEL_PATHS = [
    "/home/netview/Downloads/invertory/serialized_ONT_inventory_register (1) (7).xlsx",
    "/tmp/file_attachments/invertory/serialized_ONT_inventory_register (1) (7).xlsx",
]

def get_excel_path():
    for p in EXCEL_PATHS:
        if os.path.exists(p):
            return p
    return None

def parse_excel_date(val):
    if not val:
        return datetime.now().isoformat()
    if isinstance(val, datetime):
        return val.isoformat()
    if isinstance(val, (int, float)):
        # Excel epoch starts Dec 30 1899
        try:
            return (datetime(1899, 12, 30) + timedelta(days=val)).isoformat()
        except Exception:
            return datetime.now().isoformat()
    return str(val).strip()

def normalize_status(val):
    s = str(val or 'In Stock').strip()
    if s.lower() in ['stock out', 'issued / out', 'out', 'issued']:
        return 'Issued / Out'
    if s.lower() in ['under repair', 'repair', 'faulty']:
        return 'Under Repair'
    if s.lower() in ['decommissioned', 'retired']:
        return 'Decommissioned'
    return 'In Stock'

def migrate():
    excel_path = get_excel_path()
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Seed Admin User if not exists
    cursor.execute("SELECT count(*) FROM users WHERE email='admin@ont.co.ke'")
    if cursor.fetchone()[0] == 0:
        cursor.execute("""
            INSERT INTO users (id, email, password_hash, name, role, status, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, datetime('now'), datetime('now'))
        """, (str(uuid.uuid4()), "admin@ont.co.ke", "admin123", "System Admin", "Admin", "active"))
        print("Default admin user created: admin@ont.co.ke / admin123")

    if not excel_path:
        print(f"Excel file not found in any path: {EXCEL_PATHS}")
        conn.commit()
        conn.close()
        return

    print(f"Opening Excel workbook: {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Migrate Item Catalog
    if "Item Catalog" in wb.sheetnames:
        ws = wb["Item Catalog"]
        print("Migrating Item Catalog...")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]: continue
            sku = str(row[0]).strip()
            if not sku or sku.startswith("Formula"): continue
            asset_type = str(row[1] or 'ONT').strip()
            manufacturer = str(row[2] or '').strip()
            model = str(row[3] or '').strip()
            access_tech = str(row[4] or 'GPON').strip()
            unit_cost = float(row[5] or 0.0)
            description = str(row[6] or '').strip()
            reorder_level = int(row[7] or 5)
            
            cursor.execute("""
                INSERT OR REPLACE INTO item_catalogs (sku, asset_type, manufacturer, model, access_tech, unit_cost, description, reorder_level, tracking_type, is_active, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'SERIALIZED', 1, datetime('now'), datetime('now'))
            """, (sku, asset_type, manufacturer, model, access_tech, unit_cost, description, reorder_level))

    # Migrate Serialized Inventory
    if "Serialized Inventory" in wb.sheetnames:
        ws = wb["Serialized Inventory"]
        print("Migrating Serialized Inventory...")
        seen_serials = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]: continue
            asset_id = str(row[0]).strip()
            asset_type = str(row[1] or '').strip()
            mfg = str(row[2] or '').strip()
            model = str(row[3] or '').strip()
            tech = str(row[4] or '').strip()
            product_id = str(row[5] or '').strip()
            mac = str(row[6] or '').strip()
            sn = str(row[7] or '').strip()
            status = normalize_status(row[8])
            condition = str(row[9] or 'New').strip()
            loc = str(row[10] or 'Main Store').strip()
            custodian = str(row[11] or '').strip()
            notes = str(row[12] or '').strip()
            sku = str(row[13] or '').strip()

            # Handle blank or duplicate S/Ns safely
            if not sn:
                sn = f"UNKNOWN-{asset_id}"
            elif sn in seen_serials:
                sn = f"{sn}-{uuid.uuid4().hex[:4]}"
            seen_serials.add(sn)

            cursor.execute("""
                INSERT OR REPLACE INTO serialized_inventories 
                (asset_id, sku, asset_type, manufacturer, model, access_tech, product_id, mac, serial_number, status, condition, location, custodian, notes, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now'))
            """, (asset_id, sku, asset_type, mfg, model, tech, product_id, mac, sn, status, condition, loc, custodian, notes))

    # Migrate Customers
    if "Customer Mappings" in wb.sheetnames:
        ws = wb["Customer Mappings"]
        print("Migrating Customer Mappings...")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]: continue
            c_name = str(row[0]).strip()
            acc_no = str(row[1] or f"ACC-{uuid.uuid4().hex[:6]}").strip()
            onu_sn = str(row[2] or '').strip()
            linked_ticket = str(row[3] or '').strip()
            status = str(row[4] or 'Active').strip()

            cursor.execute("""
                INSERT OR REPLACE INTO customers (id, account_number, name, subscription_type, plot_number, location, status, assigned_onu_serial, created_at, updated_at)
                VALUES (?, ?, ?, 'PPPoE', 'Plot 12', 'Nairobi North', ?, ?, datetime('now'), datetime('now'))
            """, (str(uuid.uuid4()), acc_no, c_name, status, onu_sn))

    conn.commit()
    conn.close()
    print("Migration completed successfully.")

if __name__ == "__main__":
    migrate()
